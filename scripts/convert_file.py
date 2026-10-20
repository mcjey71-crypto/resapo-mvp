"""Private conversion worker. Input/output paths are created by ConversionService."""
from pathlib import Path
import json
import math
import resource
import subprocess
import sys
import warnings
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from conversion import AUDIO_TARGETS, MAX_OUTPUT, ConversionError, category

MEDIA_FORMATS = 'mp3,wav,mov,matroska,webm,ogg,flac,aac,asf,aiff,amr,ac3,avi,mpeg,mpegts'
MAX_CHARACTERS = 100_000


def command(args, timeout=150):
    result = subprocess.run(args, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, timeout=timeout)
    if result.returncode:
        raise ConversionError('No se pudo leer o convertir este archivo. Comprueba que no esté dañado y que sus pistas sean compatibles.')
    return result.stdout


def media(source, output):
    probe = json.loads(command([
        'ffprobe', '-v', 'error', '-protocol_whitelist', 'file,pipe', '-format_whitelist', MEDIA_FORMATS,
        '-show_entries', 'format=duration:stream=codec_type,width,height,duration,disposition',
        '-of', 'json', str(source),
    ], timeout=20))
    streams = probe.get('streams', [])
    target = output.suffix[1:]
    audio_only = target in AUDIO_TARGETS
    audios = [s for s in streams if s['codec_type'] == 'audio']
    videos = [s for s in streams if s['codec_type'] == 'video']
    selected = audios[:1] if audio_only else videos[:1] + audios[:1]
    container_duration = probe.get('format', {}).get('duration')
    durations = [s.get('duration', container_duration) for s in selected]
    durations = [container_duration if d in (None, 'N/A') else d for d in durations]
    durations = [float(d) for d in durations if d not in (None, 'N/A')]
    if selected and (not durations or not all(math.isfinite(d) and d > 0 for d in durations)):
        raise ConversionError('No pudimos determinar la duración. Prueba un archivo con duración válida.')
    duration = max(durations, default=0)
    limit = 900 if audio_only else 300
    if duration > limit:
        raise ConversionError(f'Esta conversión admite hasta {limit // 60} minutos. Recorta el archivo antes de convertirlo.')
    if audio_only and not audios:
        raise ConversionError('El archivo no contiene una pista de audio que se pueda extraer.')
    if not audio_only and not videos:
        raise ConversionError('El archivo no contiene una pista de video.')
    if not audio_only and any(s.get('width', 0) * s.get('height', 0) > 3840 * 2160 for s in videos):
        raise ConversionError('La resolución de entrada supera 4K. Reduce el tamaño del video primero.')
    args = ['ffmpeg', '-hide_banner', '-loglevel', 'error', '-nostdin', '-y', '-threads', '1',
            '-protocol_whitelist', 'file,pipe', '-format_whitelist', MEDIA_FORMATS, '-i', str(source),
            '-map_metadata', '-1', '-map_chapters', '-1', '-sn', '-dn', '-threads', '2',
            '-t', str(limit + 1), '-fs', str(MAX_OUTPUT)]
    if audio_only:
        args += ['-map', '0:a:0', '-vn', '-ac', '2']
        args += {
            'mp3': ['-c:a', 'libmp3lame', '-b:a', '192k', '-ar', '44100'],
            'm4a': ['-c:a', 'aac', '-b:a', '192k', '-ar', '44100', '-movflags', '+faststart'],
            'wav': ['-c:a', 'pcm_s16le', '-ar', '44100'],
            'flac': ['-c:a', 'flac', '-ar', '44100'],
            'ogg': ['-c:a', 'libvorbis', '-q:a', '5', '-ar', '44100'],
            'opus': ['-c:a', 'libopus', '-b:a', '128k', '-ar', '48000'],
        }[target]
    else:
        pixel_aspect = 'if(gt(sar,0),sar,1)'
        ratio = f'min(1,min(1280/(iw*{pixel_aspect}),720/ih))'
        scale = (f"scale=w='max(2,trunc(iw*{pixel_aspect}*{ratio}/2)*2)':"
                 f"h='max(2,trunc(ih*{ratio}/2)*2)',setsar=1")
        args += ['-map', '0:v:0', '-map', '0:a:0?', '-filter_threads', '1', '-vf',
                 scale,
                 '-r', '30', '-pix_fmt', 'yuv420p', '-ac', '2']
        if target == 'webm':
            args += ['-c:v', 'libvpx-vp9', '-deadline', 'realtime', '-cpu-used', '6',
                     '-b:v', '0', '-crf', '35', '-c:a', 'libopus', '-b:a', '128k']
        else:
            args += ['-c:v', 'libx264', '-preset', 'veryfast', '-crf', '24', '-c:a', 'aac', '-b:a', '128k']
            if target == 'mp4':
                args += ['-movflags', '+faststart']
    command([*args, str(output)])
    if output.stat().st_size >= MAX_OUTPUT:
        raise ConversionError('El resultado supera 50 MB. Elige un formato comprimido o un archivo más corto.')
    # Do not silently return a truncated output when a resource limit intervenes.
    check = json.loads(command(['ffprobe', '-v', 'error', '-show_entries', 'format=duration', '-of', 'json', str(output)], timeout=15))
    final_duration = float(check.get('format', {}).get('duration', 0))
    known_audio_duration = bool(audios and audios[0].get('duration') not in (None, 'N/A'))
    if (not audio_only or known_audio_duration) and abs(final_duration - duration) > max(1.5, duration * .01):
        raise ConversionError('No se pudo conservar la duración completa. Prueba otro formato de salida.')
    return ('Se convirtió la primera pista de audio a estéreo. Se omiten otras pistas y metadatos.'
            if audio_only else 'Video hasta 720p y 30 fps, sin ampliar el original. Incluye la primera pista de audio, si existe. Se omiten otras pistas, subtítulos y metadatos.')


def image(source, output):
    from PIL import Image, ImageOps
    Image.MAX_IMAGE_PIXELS = 16_000_000
    warnings.simplefilter('error', Image.DecompressionBombWarning)
    with Image.open(source) as original:
        if original.width * original.height > 16_000_000:
            raise ConversionError('La imagen supera los 16 megapíxeles. Reduce su tamaño primero.')
        if getattr(original, 'n_frames', 1) != 1:
            raise ConversionError('Esta versión convierte imágenes estáticas de una sola página. No admite animaciones ni TIFF multipágina.')
        if original.format not in {'PNG', 'JPEG', 'WEBP', 'BMP', 'TIFF', 'GIF'}:
            raise ConversionError('El contenido no corresponde a una imagen compatible.')
        raster = ImageOps.exif_transpose(original).convert('RGBA')
        target = output.suffix[1:]
        if target in {'jpg', 'bmp', 'pdf'}:
            background = Image.new('RGB', raster.size, 'white')
            background.paste(raster, mask=raster.getchannel('A'))
            raster = background
        options = {'jpg': {'quality': 92}, 'webp': {'quality': 90}, 'pdf': {'resolution': 96},
                   'tiff': {'compression': 'tiff_deflate'}}.get(target, {})
        raster.save(output, format={'jpg': 'JPEG', 'tiff': 'TIFF'}.get(target, target.upper()), **options)
    return ('La transparencia se reemplazó por blanco. Se conservan las dimensiones y se aplica la orientación de la imagen.'
            if target in {'jpg', 'bmp', 'pdf'} else 'Se conservaron dimensiones y transparencia compatible. Se aplica la orientación; se omiten metadatos.')


def read_text(source):
    extension = source.suffix.lower()
    if extension in {'.txt', '.md'}:
        data = source.read_bytes()
        try:
            text = data.decode('utf-16' if data.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig')
        except UnicodeError:
            raise ConversionError('Guarda el documento de texto con codificación UTF-8 y vuelve a intentarlo.') from None
    elif extension == '.docx':
        from docx import Document
        from docx.text.paragraph import Paragraph
        with zipfile.ZipFile(source) as archive:
            if len(archive.infolist()) > 1000 or sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024:
                raise ConversionError('El documento descomprimido es demasiado grande (máximo 20 MB).')
        document = Document(source)
        text = '\n'.join(
            block.text if isinstance(block, Paragraph) else
            '\n'.join(' | '.join(cell.text for cell in row.cells) for row in block.rows)
            for block in document.iter_inner_content()
        )
    else:
        from pypdf import PdfReader
        document = PdfReader(source, strict=True)
        if document.is_encrypted:
            raise ConversionError('Usa un PDF sin contraseña.')
        if len(document.pages) > 100:
            raise ConversionError('Esta conversión admite PDF de hasta 100 páginas.')
        parts, size = [], 0
        for page in document.pages:
            part = page.extract_text() or ''
            size += len(part)
            if size > MAX_CHARACTERS:
                raise ConversionError('El documento supera los 100 000 caracteres.')
            parts.append(part)
        text = '\n\n'.join(parts)
    text = text.replace('\r\n', '\n').replace('\r', '\n').replace('\t', '    ')
    if any(ord(char) < 32 and char not in '\n\f' for char in text):
        raise ConversionError('El archivo contiene datos binarios o caracteres de control no compatibles.')
    text = text.replace('\f', '\n\n')
    if not text.strip():
        raise ConversionError('No se encontró texto. Los PDF escaneados necesitan OCR, que aún no está incluido.')
    if len(text) > MAX_CHARACTERS:
        raise ConversionError('El documento supera los 100 000 caracteres.')
    return text


def document(source, output):
    text = read_text(source)
    target = output.suffix[1:]
    if target == 'txt':
        output.write_text(text, encoding='utf-8')
    elif target == 'docx':
        from docx import Document
        doc = Document()
        for line in text.split('\n'):
            doc.add_paragraph(line)
        doc.save(output)
    else:
        from fpdf import FPDF
        from fontTools.ttLib import TTFont
        font = ROOT / 'assets/fonts/DejaVuSans.ttf'
        with TTFont(font) as face:
            glyphs = face.getBestCmap()
        if any(ord(char) not in glyphs for char in text if char != '\n'):
            raise ConversionError('El texto contiene caracteres que la fuente PDF no puede representar. Prueba TXT o DOCX.')
        pdf = FPDF()
        pdf.set_margins(18, 18, 18)
        pdf.set_auto_page_break(True, 18)
        pdf.add_font('Document', fname=font)
        pdf.set_font('Document', size=11)
        pdf.add_page()
        for line in text.split('\n'):
            pdf.multi_cell(0, 6, line or ' ', new_x='LMARGIN', new_y='NEXT')
        pdf.output(output)
    return 'Se convirtió el texto a un documento nuevo. No se conserva el diseño, las imágenes, los estilos ni las firmas del original. Markdown se trata como texto.'


def main():
    # Inherit these limits in FFmpeg too. All workers are disposable.
    resource.setrlimit(resource.RLIMIT_AS, (640 * 1024 * 1024,) * 2)
    resource.setrlimit(resource.RLIMIT_STACK, (2 * 1024 * 1024,) * 2)
    resource.setrlimit(resource.RLIMIT_CPU, (150, 155))
    resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_OUTPUT + 1024 * 1024,) * 2)
    source, output, report = map(Path, sys.argv[1:])
    try:
        kind = category(source.name)
        note = {'audio': media, 'video': media, 'image': image, 'document': document}[kind](source, output)
        if not output.is_file() or not 0 < output.stat().st_size < MAX_OUTPUT:
            raise ConversionError('El resultado está vacío o supera el límite de 50 MB.')
        report.write_text(json.dumps({'note': note}))
    except Exception as error:
        message = str(error) if isinstance(error, ConversionError) else 'El archivo está dañado, no es compatible o supera los recursos disponibles.'
        report.write_text(json.dumps({'error': message}))
        return 1
    return 0


if __name__ == '__main__':
    sys.exit(main())
